"""投诉列表重设计后端行为测试：批量导出 ids、转办更新处理人、limit 全量返回。

对应 demo/list_demo.html 原型落地（2026-08-22）：
- GET /api/tickets/export?ids=1,2 只导出指定工单
- PUT /api/tickets/<id> {handler_name} 批量转办
- GET /api/tickets?limit=200 不再被默认 limit=10 截断
"""
import io
import tempfile
from pathlib import Path

import pytest
from conftest import _autologin_admin  # noqa: F401

# 必须在导入 app 之前把数据库指向临时文件，避免测试触碰真实数据
_TMP_DIR = Path(tempfile.mkdtemp(prefix="list-batch-tests-"))

import database  # noqa: E402

database.DB_PATH = _TMP_DIR / "test.db"
database._local.clear()

import app as app_module  # noqa: E402


@pytest.fixture()
def fresh_db(tmp_path):
    database.DB_PATH = tmp_path / "test.db"
    database._local.clear()
    database.init_db()
    yield database
    database._local.clear()


@pytest.fixture()
def client(fresh_db):
    app_module.app.config["TESTING"] = True
    with app_module.app.test_client() as c:
        try: _autologin_admin(c)
        except Exception: pass
        yield c


def _make_ticket(**overrides):
    data = {
        "student_name": "张三",
        "id_card": "110101199003070011",
        "complaint_date": "2026-05-01",
        "complaint_type": "A",
        "school_short": "南城",
        "phone": "13800000000",
    }
    data.update(overrides)
    return database.save_ticket(data)


def test_export_with_ids_only_exports_selected(client, fresh_db):
    t1 = _make_ticket(student_name="甲")
    t2 = _make_ticket(student_name="乙", id_card="110101199003070011")
    _make_ticket(student_name="丙", id_card="110101199003070011")

    resp = client.get(f"/api/tickets/export?ids={t1},{t2}")
    assert resp.status_code == 200

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    names = {r[2] for r in rows}
    assert names == {"甲", "乙"}


def test_export_without_ids_exports_all(client, fresh_db):
    _make_ticket(student_name="甲")
    _make_ticket(student_name="乙", id_card="110101199003070011")

    resp = client.get("/api/tickets/export")
    assert resp.status_code == 200

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) >= 2


def test_transfer_updates_handler(client, fresh_db):
    t = _make_ticket()
    resp = client.put(f"/api/tickets/{t}", json={"handler_name": "李老师"})
    assert resp.status_code == 200

    detail = client.get(f"/api/tickets/{t}").get_json()
    assert detail["data"]["handler_name"] == "李老师"


def test_transfer_handler_not_blocked_by_archive_gate(client, fresh_db):
    """仅改处理人不触发完结三闸门，也不需要费用解锁。"""
    t = _make_ticket()
    resp = client.put(f"/api/tickets/{t}", json={"handler_name": "王老师"})
    assert resp.status_code == 200


def test_tickets_limit_200_returns_all(client, fresh_db):
    for i in range(12):
        _make_ticket(
            student_name=f"学员{i}",
            id_card=f"44190019900101{i:04d}",
            complaint_date=f"2026-08-{i + 1:02d}",
        )
    resp = client.get("/api/tickets?limit=200").get_json()
    assert resp["data"]["total"] == 12
    assert len(resp["data"]["records"]) == 12


def test_tickets_default_limit_still_paginates(client, fresh_db):
    """默认行为不变：不带 limit 时仍按后端默认分页。"""
    for i in range(12):
        _make_ticket(
            student_name=f"学员{i}",
            id_card=f"44190019900201{i:04d}",
            complaint_date=f"2026-07-{i + 1:02d}",
        )
    resp = client.get("/api/tickets").get_json()
    assert resp["data"]["total"] == 12
    assert len(resp["data"]["records"]) <= 10


def test_delete_ticket_removes_record(client, fresh_db):
    t = _make_ticket()
    resp = client.delete(f"/api/tickets/{t}")
    assert resp.status_code == 200

    detail = client.get(f"/api/tickets/{t}")
    assert detail.status_code == 404
    assert database.get_ticket(t) is None


def test_delete_ticket_missing_returns_404(client, fresh_db):
    resp = client.delete("/api/tickets/NO-SUCH-ID")
    assert resp.status_code == 404
