"""登记表生成 API 集成测试：/api/tickets/<id>/register-form body 覆盖语义。

覆盖：
1. body 带 handling_notes/student_name → docx 内容与库内值同步覆盖（生成即预览所见）
2. body 为空 → 沿用库内值（向后兼容）
"""
import tempfile
from pathlib import Path

import pytest
from docx import Document

_TMP_DIR = Path(tempfile.mkdtemp(prefix="register-form-tests-"))

import database  # noqa: E402

database.DB_PATH = _TMP_DIR / "test.db"
database._local.clear()

import app as app_module  # noqa: E402
import config as config_module  # noqa: E402


@pytest.fixture()
def fresh_db():
    database._local.clear()
    database.init_db()
    yield database
    database._local.clear()


@pytest.fixture()
def client(fresh_db):
    app_module.app.config["TESTING"] = True
    with app_module.app.test_client() as c:
        yield c


@pytest.fixture(autouse=True)
def _reply_dir(tmp_path, monkeypatch):
    # 只在 visit_service 内替换 load_config，不写真实 data/config.json
    import services.visit_service as visit_service_module

    real = config_module.load_config()

    def _fake_load():
        cfg = real.copy()
        cfg["paths"] = dict(real.get("paths", {}), reply_dir=str(tmp_path))
        return cfg

    monkeypatch.setattr(visit_service_module, "load_config", _fake_load)
    yield tmp_path


def _make_ticket(**overrides):
    data = {
        "student_name": "张三",
        "id_card": "110101199003070011",
        "complaint_date": "2026-08-22",
        "complaint_type": "A",
        "school_short": "南城",
        "handling_notes": "库内旧处理情况",
        "fee_plan_status": "confirmed",
        "total_fee": 3700,
        "actual_paid": 3700,
    }
    data.update(overrides)
    tid = database.save_ticket(data)
    return database.get_ticket(tid)


def _docx_text(filepath):
    doc = Document(filepath)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_register_form_body_override(client, fresh_db):
    """body 覆盖 handling_notes/student_name：docx 与 DB 同步，生成即预览所见"""
    ticket = _make_ticket()
    resp = client.post(f"/api/tickets/{ticket['id']}/register-form", json={
        "handling_notes": "页面新编辑的处理情况",
        "student_name": "张三补",
    })
    body = resp.get_json()
    assert body["success"] is True, body
    assert Path(body["data"]["filepath"]).is_file()

    text = _docx_text(body["data"]["filepath"])
    assert "页面新编辑的处理情况" in text
    assert "库内旧处理情况" not in text

    row = database.get_ticket(ticket["id"])
    assert row["handling_notes"] == "页面新编辑的处理情况"
    assert row["student_name"] == "张三补"
    assert row["registration_form_path"] == body["data"]["filepath"]


def test_register_form_no_body_uses_db(client, fresh_db):
    """body 为空（旧客户端/归档链路）→ 沿用库内值，行为不变"""
    ticket = _make_ticket()
    resp = client.post(f"/api/tickets/{ticket['id']}/register-form", json={})
    body = resp.get_json()
    assert body["success"] is True, body

    text = _docx_text(body["data"]["filepath"])
    assert "库内旧处理情况" in text

    row = database.get_ticket(ticket["id"])
    assert row["handling_notes"] == "库内旧处理情况"
    assert row["student_name"] == "张三"
