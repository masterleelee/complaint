"""Word（.docx）合同上传支持测试（P1-2）。

背景：此前上传校验 CONTRACT_UPLOAD_EXTENSIONS 缺 .docx，且文本提取链路
extract_contract_text_from_file 只有 pdf/图片分支，docx 上传会「格式不支持」或
「无法提取文本」。本轮打通 docx → 文本提取整条链路。

覆盖：
1. file_parser._extract_docx 遍历表格（费用表等关键字段常落在 docx 表格里，漏掉致缺失）。
2. extract_contract_text_from_file 对 .docx 返回 source=docx_text 的文本。
3. app.CONTRACT_UPLOAD_EXTENSIONS 含 .docx。
"""
import tempfile
from pathlib import Path

import pytest

_TMP_DIR = Path(tempfile.mkdtemp(prefix="docx-upload-tests-"))

from docx import Document  # noqa: E402


def _make_docx(path: Path, paragraphs: list[str], table_rows: list[list[str]] | None = None) -> str:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, cell_text in enumerate(row):
                table.rows[i].cells[j].text = cell_text
    doc.save(str(path))
    return str(path)


# ── 1. file_parser._extract_docx 遍历表格 ──────────────────────────────

def test_extract_docx_includes_table_text(tmp_path):
    from services.file_parser import _extract_docx

    path = _make_docx(
        tmp_path / "with_table.docx",
        paragraphs=["第六条 退学退费", "已发生的实际操作培训费如下表"],
        table_rows=[["科目二实操费", "80元/学时"], ["科目三实操费", "80元/学时"]],
    )
    text = _extract_docx(path)
    # 段落文本保留
    assert "第六条 退学退费" in text
    # 表格文本（费用表）被提取出来
    assert "科目二实操费" in text
    assert "80元/学时" in text
    assert "科目三实操费" in text


def test_extract_docx_no_table_still_works(tmp_path):
    from services.file_parser import _extract_docx

    path = _make_docx(tmp_path / "plain.docx", paragraphs=["第一条 总则", "第二条 服务内容"])
    text = _extract_docx(path)
    assert "第一条 总则" in text
    assert "第二条 服务内容" in text


# ── 2. extract_contract_text_from_file 支持 .docx ──────────────────────

def test_extract_contract_text_from_file_docx(tmp_path):
    from services.contract_service import extract_contract_text_from_file

    path = _make_docx(
        tmp_path / "contract.docx",
        paragraphs=[
            "东莞市机动车驾驶员培训合同",
            "第五条 退学退费",
            "违约金为全部培训费用的20%",
            "已发生的实际操作培训费＝已产生的实际操作培训学时×约定的学时收费标准是80元/学时",
        ],
        table_rows=[["费用项目", "金额"], ["服务费", "1000"]],
    )
    result = extract_contract_text_from_file(path)
    assert "error" not in result
    assert result["source"] == "docx_text"
    assert "80元/学时" in result["text"]
    # 表格文本也进来了（服务费金额）
    assert "服务费" in result["text"]


def test_extract_contract_text_from_file_missing_docx(tmp_path):
    from services.contract_service import extract_contract_text_from_file

    result = extract_contract_text_from_file(str(tmp_path / "不存在.docx"))
    assert result.get("error") == "合同文件不存在"


# ── 3. app.CONTRACT_UPLOAD_EXTENSIONS 含 .docx ─────────────────────────

def test_upload_extensions_include_docx():
    import app as app_module  # noqa: F401

    assert ".docx" in app_module.CONTRACT_UPLOAD_EXTENSIONS
    assert ".pdf" in app_module.CONTRACT_UPLOAD_EXTENSIONS
