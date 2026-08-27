"""数据质量与统计口径修复测试。

覆盖：
1. 统计三桶排除已撤诉 + 新增撤诉独立桶（合计=总投诉数）+ by_source/趋势排除撤诉
2. 超期口径与前端 isOverdue 对齐（投诉日期超7天、未完结、未撤诉未归档）
3. 导出：新增撤诉状态/归档状态两列、scope 网点过滤、ids 直查、公式注入转义
4. save_ticket complaint_date 非法格式回退当天
5. 同日同人去重保留原投诉内容与来源渠道
6. 手机号服务端格式校验 400
7. 查询双空结果不建工单且错误文案透传轮询接口
"""
import io
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

import database  # noqa: E402


@pytest.fixture()
def fresh_db(tmp_path):
    database.DB_PATH = tmp_path / "test.db"
    database._local.clear()
    database.init_db()
    yield database
    database._local.clear()


@pytest.fixture()
def client(fresh_db):
    from app import app as flask_app

    flask_app.config["TESTING"] = True
    with flask_app.test_client() as c:
        yield c


def _make_ticket(**overrides):
    data = {
        "student_name": "张三",
        "id_card": "110101199003070011",
        "complaint_date": "2026-08-20",
        "complaint_type": "A",
        "school_short": "南城",
    }
    data.update(overrides)
    return database.save_ticket(data)


def test_stats_withdrawn_excluded_from_buckets_and_new_bucket(fresh_db):
    active_id = _make_ticket(
        id_card="110101199003070011",
        complaint_date="2031-01-02",
        source_channel="有效渠道X",
        handle_status="待处理",
    )
    withdrawn_id = _make_ticket(
        id_card="110101199003070011",
        complaint_date="2031-01-01",
        source_channel="撤诉专属渠道",
        handle_status="待处理",
    )
    database.update_ticket(withdrawn_id, {"withdraw_status": "已撤诉"})

    stats = database.get_ticket_statistics()
    assert stats["total"] == 2
    assert stats["pending"] == 1
    assert stats["withdrawn"] == 1
    assert stats["pending"] + stats["processing"] + stats["completed"] + stats["withdrawn"] == stats["total"]
    assert stats["effective_total"] == 1

    by_source = {item["source"]: item["count"] for item in stats["by_source"]}
    assert by_source.get("有效渠道X") == 1
    assert "撤诉专属渠道" not in by_source

    daily = {row["date"]: row["count"] for row in stats["daily_trend"]}
    assert daily.get("2031-01-02") == 1
    assert "2031-01-01" not in daily

    monthly = {row["month"]: row["count"] for row in stats["monthly_trend"]}
    assert monthly.get("2031-01") == 1


def test_overdue_uses_seven_day_frontend_criteria(fresh_db):
    old = (date.today() - timedelta(days=10)).isoformat()
    recent = (date.today() - timedelta(days=5)).isoformat()

    _make_ticket(id_card="110101199003070011", complaint_date=old, handle_status="处理中")
    _make_ticket(id_card="110101199003070011", complaint_date=recent, handle_status="处理中")
    _make_ticket(
        id_card="110101199003070011", complaint_date=old,
        handle_status="待处理", withdraw_status="已撤诉",
    )
    _make_ticket(
        id_card="110101199003070011", complaint_date=old,
        handle_status="待处理", archive_status="已归档",
    )
    _make_ticket(id_card="110101199003070011", complaint_date=old, handle_status="已完结")

    stats = database.get_processing_duration_stats()
    assert stats["overdue_count"] == 1


def _load_export_rows(resp_or_bytes):
    from openpyxl import load_workbook

    payload = resp_or_bytes if isinstance(resp_or_bytes, bytes) else resp_or_bytes.data
    wb = load_workbook(io.BytesIO(payload))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return headers, rows


def test_export_new_columns_scope_ids_and_formula_escape(client):
    branch_id = _make_ticket(
        student_name="=SUM(A1)",
        id_card="110101199003070011",
        organization_unit_type="分校",
        organization_unit_name="南城分校",
        withdraw_status="已撤诉",
    )
    store_id = _make_ticket(
        student_name="@风险姓名",
        id_card="110101199003070011",
        organization_unit_type="分店",
        organization_unit_name="东城分店",
        archive_status="已归档",
    )

    headers, rows = _load_export_rows(client.get("/api/tickets/export"))
    assert headers[-2:] == ["撤诉状态", "归档状态"]

    _, scoped_rows = _load_export_rows(client.get("/api/tickets/export?scope=branch").data)
    assert {r[0] for r in scoped_rows} == {branch_id}
    _, store_rows = _load_export_rows(client.get("/api/tickets/export?scope=store").data)
    assert {r[0] for r in store_rows} == {store_id}

    resp = client.get(f"/api/tickets/export?ids={branch_id}&status=不存在的状态")
    assert resp.status_code == 200
    _, ids_rows = _load_export_rows(resp)
    assert {r[0] for r in ids_rows} == {branch_id}

    by_name = {r[0]: r for r in rows}
    assert by_name[branch_id][2].startswith("'=")
    assert by_name[branch_id][-2] == "已撤诉"
    assert not by_name[branch_id][-1]
    assert by_name[store_id][2].startswith("'@")
    assert by_name[store_id][-1] == "已归档"

    assert client.get("/api/tickets/export?scope=hacker").status_code == 400


def test_save_ticket_normalizes_invalid_complaint_date(fresh_db):
    today = datetime.now().strftime("%Y-%m-%d")
    t1 = _make_ticket(id_card="110101199003070011", complaint_date="2026/08/23")
    assert database.get_ticket(t1)["complaint_date"] == today
    t2 = _make_ticket(id_card="110101199003070011", complaint_date="2026-13-99")
    assert database.get_ticket(t2)["complaint_date"] == today
    t3 = _make_ticket(id_card="110101199003070011", complaint_date="")
    assert database.get_ticket(t3)["complaint_date"] == today
    t4 = _make_ticket(id_card="110101199003070011", complaint_date="2026-08-23")
    assert database.get_ticket(t4)["complaint_date"] == "2026-08-23"


def test_dedup_preserves_original_content_and_channel(fresh_db):
    id_card = "110101199003070011"
    first = _make_ticket(
        id_card=id_card,
        complaint_date="2026-08-20",
        complaint_content="原始投诉内容",
        source_channel="交通部门",
    )
    second = _make_ticket(
        id_card=id_card,
        complaint_date="2026-08-20",
        complaint_content="二次提交内容",
        source_channel="12345热线",
        handler_name="新处理人",
    )
    assert second == first
    ticket = database.get_ticket(first)
    assert ticket["complaint_content"] == "原始投诉内容"
    assert ticket["source_channel"] == "交通部门"
    assert ticket["handler_name"] == "新处理人"


def test_invalid_phone_rejected_with_400(client):
    resp = client.post("/api/query", json={"phone": "12345678901"})
    assert resp.status_code == 400
    assert "手机号" in resp.get_json()["error"]

    resp_start = client.post("/api/query/start", json={"phone": "12345"})
    assert resp_start.status_code == 400
    assert "手机号" in resp_start.get_json()["error"]


def test_empty_query_result_creates_no_ticket(client, fresh_db, monkeypatch):
    app_module = __import__("app", fromlist=["app"])
    job_id = "test-job-empty-result"
    with app_module.QUERY_JOBS_LOCK:
        app_module.QUERY_JOBS[job_id] = {
            "id": job_id,
            "status": "queued",
            "created_at": "",
            "result": None,
            "sources": {"internal": "pending", "third": "pending", "driving": "pending"},
            "query_durations_ms": {},
            "error": "",
        }

    def fake_query_by_phone(phone, timeout=60.0):
        return {
            "name": "", "id_card": "", "phone": phone,
            "sources": {"internal": "not_found", "third": "not_found", "driving": "not_found"},
        }

    monkeypatch.setattr(app_module, "query_all_systems_by_phone", fake_query_by_phone)
    app_module._query_job_worker(job_id, {"phone": "13800000000"})

    job = app_module.QUERY_JOBS[job_id]
    assert job["status"] == "failed"
    assert job["error"] == "未匹配到学员档案，请核对手机号或改用身份证号查询"

    records, _ = database.list_tickets(limit=100)
    assert all(r["phone"] != "13800000000" for r in records)

    polled = client.get(f"/api/query/status/{job_id}").get_json()
    assert polled["success"] is True
    assert polled["error"] == job["error"]
