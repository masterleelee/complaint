"""归档 API 集成测试（实施Agent K）：/api/reply/generate (v2) 与 /api/tickets/<id>/archive。

覆盖：
1. reply/generate 返回成功且生成 docx（标题居中、落款右对齐）
2. archive 三闸门缺失返回 400 并含 errors
3. archive 闸门全过 + 预置两件套 → 成功返回且目标文件存在
"""
import tempfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

_TMP_DIR = Path(tempfile.mkdtemp(prefix="archive-api-tests-"))

import database  # noqa: E402

database.DB_PATH = _TMP_DIR / "test.db"
database._local.clear()

import app as app_module  # noqa: E402
import config as config_module  # noqa: E402
import services.archive_service as archive_service_module  # noqa: E402
from services.reply_docx import generate_reply_docx  # noqa: E402


@pytest.fixture()
def fresh_db():
    database._local.clear()
    database.init_db()
    yield database
    database._local.clear()


@pytest.fixture()
def client(fresh_db):
    app_module.app.config["TESTING"] = True
    with app_module.app.test_client() as c:
        yield c


def _make_ticket(**overrides):
    data = {
        "student_name": "张三",
        "id_card": "110101199003070011",
        "complaint_date": "2026-08-22",
        "complaint_type": "A",
        "school_short": "南城",
        "organization_unit_type": "分校",
        "organization_unit_name": "南城分校",
        "registration_date": "2026-01-01",
        "license_type": "C2",
        "exam_stage": "科目二",
        "total_fee": 3700,
        "actual_paid": 3700,
        "fee_plan_status": "confirmed",
    }
    data.update(overrides)
    tid = database.save_ticket(data)
    return database.get_ticket(tid)


# ───────────────────────────────────────────────────────────
# 1) /api/reply/generate → v2 docx 标题居中 + 落款右对齐
# ───────────────────────────────────────────────────────────
def test_reply_generate_v2_format(client, fresh_db, tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "get_archive_folder", lambda *a, **k: str(tmp_path))
    ticket = _make_ticket()

    resp = client.post("/api/reply/generate", json={
        "ticket_id": ticket["id"],
        "deductions": [
            {"item": "综合服务费扣除", "amt": 1100, "basis": "合同第七条"},
            {"item": "理论培训费扣除", "amt": 600, "basis": "合同第七条"},
            {"item": "已退费用", "amt": 0, "basis": "应跳过"},  # <=0 跳过
        ],
    })
    body = resp.get_json()
    assert body["success"] is True
    assert body["filepath"].endswith(".docx")

    doc = Document(body["filepath"])
    # 标题：首个段落居中
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "回复" in doc.paragraphs[0].text

    # 落款：存在右对齐段落
    right_count = sum(1 for p in doc.paragraphs if p.alignment == WD_ALIGN_PARAGRAPH.RIGHT)
    assert right_count >= 1


# ───────────────────────────────────────────────────────────
# 2) archive 三闸门缺失 → 400 + errors
# ───────────────────────────────────────────────────────────
def test_archive_gate_missing_returns_400(client, fresh_db, tmp_path, monkeypatch):
    monkeypatch.setattr(config_module, "load_config", lambda: {"archive_root": str(tmp_path)})
    monkeypatch.setattr(archive_service_module, "load_config", lambda: {"archive_root": str(tmp_path)})
    ticket = _make_ticket(handling_notes="", branch_cooperation="", fee_plan_status="draft")

    resp = client.post(f"/api/tickets/{ticket['id']}/archive")
    assert resp.status_code == 400
    body = resp.get_json()
    assert body["success"] is False
    assert isinstance(body["errors"], list)
    assert "处理情况未填写" in body["errors"]
    assert "配合度未评定" in body["errors"]
    assert "费用明细未确认" in body["errors"]


# ───────────────────────────────────────────────────────────
# 3) archive 闸门全过 + 预置两件套 → 成功且目标文件存在
# ───────────────────────────────────────────────────────────
def test_archive_success_copies_two_files(client, fresh_db, tmp_path, monkeypatch):
    monkeypatch.setattr(config_module, "load_config", lambda: {"archive_root": str(tmp_path)})
    monkeypatch.setattr(archive_service_module, "load_config", lambda: {"archive_root": str(tmp_path)})
    ticket = _make_ticket(
        handling_notes="已协调网点退费",
        branch_cooperation="配合",
        fee_plan_status="confirmed",
    )

    # 预置登记表 / 回复函临时文件（沿用受理写盘字段命名）
    reg_src = tmp_path / "登记表_src.docx"
    reply_src = tmp_path / "回复函_src.docx"
    Document().save(str(reg_src))
    Document().save(str(reply_src))
    database.update_ticket(ticket["id"], {
        "registration_form_path": str(reg_src),
        "reply_path": str(reply_src),
    })

    resp = client.post(f"/api/tickets/{ticket['id']}/archive")
    body = resp.get_json()
    assert body["success"] is True
    assert body["opened"] is False
    assert len(body["files"]) == 2

    # 目标文件按 build_archive_dir 规则落盘
    case_dir, reg_target, reply_target = app_module.build_archive_dir(
        database.get_ticket(ticket["id"]), root=str(tmp_path))
    assert Path(reg_target).is_file()
    assert Path(reply_target).is_file()

    # 工单状态更新
    assert database.get_ticket(ticket["id"])["archive_status"] == "已归档"


# ───────────────────────────────────────────────────────────
# 4) archive 无两件套 → 400 拒绝
# ───────────────────────────────────────────────────────────
def test_archive_no_files_returns_400(client, fresh_db, tmp_path, monkeypatch):
    monkeypatch.setattr(config_module, "load_config", lambda: {"archive_root": str(tmp_path)})
    monkeypatch.setattr(archive_service_module, "load_config", lambda: {"archive_root": str(tmp_path)})
    # 服务端重生成登记表失败（模拟），确保走到「无可归档文件」分支
    monkeypatch.setattr(app_module, "_regen_registration_form",
                        lambda t, output_dir="": {"success": False, "error": "regen unavailable"})
    ticket = _make_ticket(
        handling_notes="已协调网点退费",
        branch_cooperation="配合",
        fee_plan_status="confirmed",
    )
    # 不预置任何源文件
    database.update_ticket(ticket["id"], {
        "registration_form_path": "",
        "reply_path": "",
    })

    resp = client.post(f"/api/tickets/{ticket['id']}/archive")
    assert resp.status_code == 400
    body = resp.get_json()
    assert body["success"] is False
    assert "无可归档文件" in body["errors"][0]
