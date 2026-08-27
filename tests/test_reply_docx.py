"""回复函 docx 生成服务测试：标题文本/对齐/扣费清单行。"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from services.reply_docx import generate_reply_docx

TICKET = {
    "student_name": "陈虹妃",
    "id_card": "110101199003070011",
    "complaint_date": "2026-04-20",
    "registration_date": "2025-04-01",
    "school_name": "塘厦林村",
    "school_short": "青C",
    "organization_unit_type": "分校",
    "license_type": "C2",
    "exam_stage": "科目二",
    "total_fee": 3700,
    "actual_paid": 3700,
    "contract_code": "DGJP202504100026",
    "training_hours": {"科目一": "12时0分", "科目二": "17时2分", "科目三": "27时15分"},
}

DEDUCTIONS = [
    {"item": "综合服务费扣除", "amt": 1100, "basis": "合同第七条第一款"},
    {"item": "理论培训费扣除", "amt": 600, "basis": "合同第七条第一款"},
    {"item": "已退费用", "amt": 0, "basis": "不应出现"},      # <=0 应跳过
    {"item": "负数项", "amt": -50, "basis": "不应出现"},      # <=0 应跳过
]


def _generate(tmp_path, name="reply.docx"):
    out = str(tmp_path / name)
    result = generate_reply_docx(TICKET, DEDUCTIONS, out)
    assert result["success"], f"生成失败: {result.get('error')}"
    return result["path"], Document(out)


def test_generate_success_and_title_text(tmp_path):
    path, doc = _generate(tmp_path)
    assert path.endswith("reply.docx")
    assert doc.paragraphs[0].text == "关于陈虹妃投诉的回复"


def test_unit_type_not_duplicated_when_school_name_ends_with_it(tmp_path):
    ticket = dict(TICKET, school_name="横沥河畔分校")
    out = str(tmp_path / "dup.docx")
    result = generate_reply_docx(ticket, DEDUCTIONS, out)
    assert result["success"]
    doc = Document(out)
    body = next(p.text for p in doc.paragraphs if p.text.startswith("经我驾校调查核实"))
    assert "在横沥河畔分校网点报名C2驾照培训" in body
    assert "分校分校" not in body


def test_title_center_and_signature_right(tmp_path):
    _, doc = _generate(tmp_path)
    texts = [p.text for p in doc.paragraphs]

    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    sign_idx = texts.index("驾校名称：东莞市快捷汽车驾驶员培训有限公司（公章）")
    date_idx = next(i for i, t in enumerate(texts) if t.endswith("日") and "年" in t)
    assert doc.paragraphs[sign_idx].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert doc.paragraphs[date_idx].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_deduction_lines_match_nonzero_rows_with_amounts(tmp_path):
    _, doc = _generate(tmp_path)
    lines = [p.text for p in doc.paragraphs]
    fee_lines = [t for t in lines if t[:2] in ("1、", "2、") and "元" in t]

    nonzero = [d for d in DEDUCTIONS if d["amt"] > 0]
    assert len(fee_lines) == len(nonzero)
    assert any("1100元" in t for t in fee_lines)
    assert any("600元" in t for t in fee_lines)
    assert not any("不应出现" in t for t in lines)


def test_body_contains_progress_hours_and_contract_clause(tmp_path):
    _, doc = _generate(tmp_path)
    lines = [p.text for p in doc.paragraphs]

    known = next(t for t in lines if t.startswith("据了解，"))
    assert "学员报名共交培训服务费3700元" in known
    assert "目前进度处于：科目二阶段" in known
    assert "科目三培训27时15分、科目二培训17时2分" in known

    clause = next(t for t in lines if t.startswith("按照《东莞市机动车驾驶员培训服务合同》"))
    assert "（合同编码：DGJP202504100026）" in clause
    assert "第九条退学退费相关约定" in clause
    assert "（一）合同有效期内，甲方因个人原因中途提出退学的应向乙方提交书面申请" in clause
    assert clause.endswith("扣费如下：")

    body = next(t for t in lines if t.startswith("经我驾校调查核实"))
    assert "在塘厦林村分校网点报名C2驾照培训" in body


def test_missing_fields_tolerated_and_total_bold(tmp_path):
    out = str(tmp_path / "bare.docx")
    result = generate_reply_docx({}, [], out)
    assert result["success"], f"空字段应容错不报错: {result.get('error')}"
    doc = Document(out)
    total = next(p for p in doc.paragraphs if p.text.startswith("总扣费："))
    assert total.text == "总扣费：0元。学员实际已交费用0元，应退回：0元。"
    assert all(run.bold for run in total.runs if run.text.strip())
    lines = [p.text for p in doc.paragraphs]
    known = next(t for t in lines if t.startswith("据了解，"))
    assert known == "据了解，学员报名共交培训服务费0元，目前进度处于：未知阶段。"
    clause = next(t for t in lines if t.startswith("按照《东莞市机动车驾驶员培训服务合同》"))
    assert "合同编码" not in clause
