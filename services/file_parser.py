"""多格式文件文本提取器 - 支持 PDF/图片/DOCX/XLSX/TXT

图片 OCR = macOS 原生 Vision（系统自带，离线、亚秒级）。
EasyOCR/PaddleOCR 已按 spec（.scratch/contract-ocr-baidu/spec.md v3 §6.2）删除：
实测生产从未触发（Vision 先成功 / Paddle 未安装），删除省 ~500M 依赖并加快启动。
"""
import os
from utils.logger import system_logger


def extract_text(filepath: str) -> str:
    """根据文件扩展名自动选择提取方式"""
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(filepath)
    elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
        return _extract_image(filepath)
    elif ext == ".docx":
        return _extract_docx(filepath)
    elif ext == ".xlsx":
        return _extract_xlsx(filepath)
    elif ext == ".txt":
        return _extract_txt(filepath)
    else:
        return ""


def _extract_pdf(filepath: str) -> str:
    """从 PDF 提取文本；文字层为空时（扫描件）回退到图片 OCR"""
    text = _extract_pdf_text_layer(filepath)
    if len(text.strip()) >= 10:
        return text
    return _extract_pdf_ocr(filepath)


def _extract_pdf_text_layer(filepath: str) -> str:
    """从 PDF 文字层提取文本"""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)
    except Exception:
        pass

    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(filepath)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        return ""


def _extract_pdf_ocr(filepath: str) -> str:
    """扫描件 PDF：逐页渲染成图片后走 OCR"""
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(filepath)
        lines = []
        for i in range(len(pdf)):
            bitmap = pdf[i].render(scale=2)
            pil_image = bitmap.to_pil()
            temp_path = f"{filepath}_page{i}_ocr.png"
            pil_image.convert('RGB').save(temp_path)
            try:
                page_text = _extract_image(temp_path)
                if page_text.strip():
                    lines.append(page_text)
            finally:
                if os.path.exists(temp_path):
                    os.remove(temp_path)
        return "\n".join(lines)
    except Exception as e:
        system_logger.warning("PDF OCR Error: %s", e)
        return ""


def _extract_image(filepath: str) -> str:
    """从图片提取文本：macOS Vision（系统原生，亚秒级；离线兜底）。"""
    return _extract_image_vision(filepath)


def _extract_image_vision(filepath: str) -> str:
    """使用 macOS 系统 Vision 框架识别图片文本（需 pyobjc-framework-Vision）"""
    try:
        import Vision
        from Foundation import NSURL

        url = NSURL.fileURLWithPath_(os.path.abspath(filepath))
        handler = Vision.VNImageRequestHandler.alloc().initWithURL_options_(url, None)
        request = Vision.VNRecognizeTextRequest.alloc().init()
        request.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
        request.setRecognitionLanguages_(["zh-Hans", "en-US"])
        request.setUsesLanguageCorrection_(True)
        ok, err = handler.performRequests_error_([request], None)
        if not ok:
            raise RuntimeError(str(err))
        lines = [r.topCandidates_(1)[0].string() for r in request.results()]
        return "\n".join(line for line in lines if line)
    except ImportError:
        system_logger.info("[OCR] pyobjc-framework-Vision 未安装，无法本地识别")
        return ""
    except Exception as e:
        system_logger.warning("Vision OCR Error: %s", e)
        return ""


def warmup_ocr():
    """预热 OCR 引擎（预加载 macOS Vision 框架），消除首个图片请求的冷启动延迟。"""
    try:
        import Vision  # noqa: F401 预加载 macOS Vision 框架
    except Exception:
        pass


def _extract_docx(filepath: str) -> str:
    """从 Word 文档提取文本（含表格——费用表等关键字段常落在 docx 表格里，漏掉会致关键字段缺失）。"""
    try:
        from docx import Document
        doc = Document(filepath)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _extract_xlsx(filepath: str) -> str:
    """从 Excel 提取所有单元格文本"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"--- 工作表: {sheet_name} ---")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except Exception:
        return ""


def _extract_txt(filepath: str) -> str:
    """从文本文件提取"""
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        try:
            with open(filepath, "r", encoding="gbk") as f:
                return f.read()
        except Exception:
            return ""
